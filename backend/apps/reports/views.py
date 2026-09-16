import csv
from io import BytesIO

import openpyxl
from django.http import HttpResponse
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework import permissions
from rest_framework.views import APIView

from apps.dashboards.views import month_bounds
from apps.transactions.filters import TransactionFilter
from apps.transactions.selectors import get_user_transactions
from apps.transactions.services import compute_expenses_by_category, compute_monthly_summary

EXPORT_COLUMNS = ["Data", "Descrição", "Categoria", "Tipo", "Valor", "Situação"]


def _export_rows(request):
    queryset = get_user_transactions(user=request.user)
    queryset = TransactionFilter(request.query_params, queryset=queryset).qs
    for transaction in queryset:
        yield [
            transaction.competence_date.strftime("%d/%m/%Y"),
            transaction.title,
            transaction.category.name,
            transaction.get_transaction_type_display(),
            transaction.total_amount,
            transaction.get_status_display(),
        ]


class TransactionExportView(APIView):
    """Flexible export honoring the same filters as GET /api/transactions/
    (transaction_type, category, status, date_from, date_to) — one endpoint
    that covers monthly/annual summaries, per-category breakdowns and
    pending-amounts reports depending on how it's filtered, instead of a
    rigid endpoint per report type."""

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        # Named "type", not "format": DRF reserves the `format` query param
        # for its own content-negotiation override and 404s when it doesn't
        # match a registered renderer (see negotiation.filter_renderers).
        export_format = request.query_params.get("type", "csv")
        if export_format == "xlsx":
            return self._as_xlsx(request)
        return self._as_csv(request)

    def _as_csv(self, request):
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="transacoes.csv"'
        response.write("﻿")  # BOM so Excel opens the UTF-8 file with accents intact
        writer = csv.writer(response)
        writer.writerow(EXPORT_COLUMNS)
        writer.writerows(_export_rows(request))
        return response

    def _as_xlsx(self, request):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Transações"
        sheet.append(EXPORT_COLUMNS)
        for row in _export_rows(request):
            sheet.append(row)
        for index, header in enumerate(EXPORT_COLUMNS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(12, len(header) + 6)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="transacoes.xlsx"'
        workbook.save(response)
        return response


def _format_currency(value):
    return f"R$ {value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


class MonthlySummaryPdfView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        start, end = month_bounds(request.query_params.get("month"))
        user = request.user

        summary = compute_monthly_summary(user=user, start=start, end=end)
        expenses_by_category = compute_expenses_by_category(user=user, start=start, end=end)

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
        styles = getSampleStyleSheet()

        elements = [
            Paragraph(f"Resumo financeiro — {start:%m/%Y}", styles["Title"]),
            Paragraph(user.preferred_name or user.email, styles["Normal"]),
            Spacer(1, 0.6 * cm),
        ]

        summary_rows = [
            ["Receita recebida", _format_currency(summary["income_received"])],
            ["Receita prevista", _format_currency(summary["income_total"])],
            ["Despesa paga", _format_currency(summary["expense_paid"])],
            ["Despesa prevista", _format_currency(summary["expense_total"])],
            ["Saldo do mês (caixa)", _format_currency(summary["cash_profit"])],
            ["Resultado (competência)", _format_currency(summary["accrual_profit"])],
        ]
        summary_table = Table(summary_rows, colWidths=[9 * cm, 5 * cm])
        summary_table.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f8fafc")),
                ]
            )
        )
        elements.append(summary_table)
        elements.append(Spacer(1, 1 * cm))
        elements.append(Paragraph("Despesas por categoria", styles["Heading2"]))
        elements.append(Spacer(1, 0.2 * cm))

        category_rows = [["Categoria", "Total"]] + [
            [row["category__name"], _format_currency(row["total"])] for row in expenses_by_category
        ]
        if len(category_rows) == 1:
            category_rows.append(["Nenhuma despesa neste mês", "-"])

        category_table = Table(category_rows, colWidths=[9 * cm, 5 * cm])
        category_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ]
            )
        )
        elements.append(category_table)

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer.read(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="resumo-{start:%Y-%m}.pdf"'
        return response
